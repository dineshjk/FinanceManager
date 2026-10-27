# -*- coding: utf-8 -*-
# tests/test_drive_sync.py

"""
Unit tests for StockMan Google Drive integration and startup sequence.
"""

import os
import tempfile
import sqlite3
import pytest
from unittest.mock import MagicMock, patch

import openpyxl

from StockMan import drive_sync


def test_read_investment_parameters_nonexistent():
    """Verify default values are returned when Excel file does not exist."""
    app, cash = drive_sync.read_investment_parameters("non_existent_file.xlsx")
    assert app == drive_sync.DEFAULT_APPETITE
    assert cash == drive_sync.DEFAULT_CASH


def test_read_investment_parameters_from_excel():
    """Verify parameters are read accurately from Row 1 & Row 2."""
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp_path = tmp.name
    tmp.close()

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(1, 1, "Current Appetite")
        ws.cell(1, 2, 2500000.0)
        ws.cell(2, 1, "Cash Available for Investment")
        ws.cell(2, 2, 350000.0)
        wb.save(tmp_path)
        wb.close()

        app, cash = drive_sync.read_investment_parameters(tmp_path)
        assert app == 2500000.0
        assert cash == 350000.0
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_update_portfolio_excel_structure_and_sqlite_data():
    """Verify update_portfolio_excel writes Row 1, Row 2, Row 3 headers, and stock records."""
    db_tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
    db_path = db_tmp.name
    db_tmp.close()

    xlsx_tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    xlsx_path = xlsx_tmp.name
    xlsx_tmp.close()

    try:
        # Create a mock stocks table
        conn = sqlite3.connect(db_path)
        try:
            conn.execute("""
                CREATE TABLE stocks (
                    id_stk INTEGER PRIMARY KEY,
                    stk_code TEXT,
                    isin TEXT,
                    company_name TEXT,
                    curr_avg_price REAL,
                    buy_qty INTEGER,
                    sell_qty INTEGER
                )
            """)
            conn.execute("""
                INSERT INTO stocks VALUES
                (1, 'HAL', 'INE066F01012', 'Hindustan Aeronautics Limited', 4850.50, 100, 20),
                (2, 'TCS', 'INE467B01029', 'Tata Consultancy Services', 3500.00, 50, 50)
            """)
            conn.commit()
        finally:
            conn.close()

        # Write to Excel
        drive_sync.update_portfolio_excel(
            excel_path=xlsx_path,
            current_appetite=1800000.0,
            cash_available=150000.0,
            db_path=db_path,
        )

        # Inspect generated workbook
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        ws = wb.active

        # Check Row 1 and Row 2
        assert ws.cell(1, 1).value == "Current Appetite"
        assert ws.cell(1, 2).value == 1800000.0
        assert ws.cell(2, 1).value == "Cash Available for Investment"
        assert ws.cell(2, 2).value == 150000.0

        # Check Row 3 Headers
        expected_headers = [
            "id_stk",
            "stk_code",
            "isin",
            "company_name",
            "curr_avg_price",
            "current_qty",
            "curr_investment_amt",
        ]
        actual_headers = [ws.cell(3, c).value for c in range(1, 8)]
        assert actual_headers == expected_headers

        # Check Row 4 (HAL: qty = 100 - 20 = 80, amt = 80 * 4850.50 = 388040.0)
        assert ws.cell(4, 1).value == 1
        assert ws.cell(4, 2).value == "HAL"
        assert ws.cell(4, 3).value == "INE066F01012"
        assert ws.cell(4, 4).value == "Hindustan Aeronautics Limited"
        assert ws.cell(4, 5).value == 4850.50
        assert ws.cell(4, 6).value == 80
        assert ws.cell(4, 7).value == 388040.0

        # Check Row 5 (TCS: qty = 0, amt = 0)
        assert ws.cell(5, 1).value == 2
        assert ws.cell(5, 2).value == "TCS"
        assert ws.cell(5, 6).value == 0
        assert ws.cell(5, 7).value == 0.0

        wb.close()
    finally:
        if os.path.exists(db_path):
            os.remove(db_path)
        if os.path.exists(xlsx_path):
            os.remove(xlsx_path)


def test_get_local_pick_of_the_day_with_temp_dir():
    """Verify get_local_pick_of_the_day prioritizes latest alias and glob files."""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "Pick_of_the_day_20261023.md")
        with open(test_file, "w", encoding="utf-8") as f:
            f.write("# 🏆 Pick of the Day Content")

        res = drive_sync.get_local_pick_of_the_day(data_dir=temp_dir)
        assert res is not None
        content, fname, mtime = res
        assert "# 🏆 Pick of the Day Content" in content
        assert fname == "Pick_of_the_day_20261023.md"


def test_upload_portfolio_to_drive_mock():
    """Verify upload_portfolio_to_drive handles folder lookup and file create/update."""
    mock_service = MagicMock()
    mock_files = mock_service.files.return_value

    # 1. Folder search returns Finance folder id 'folder_123'
    mock_files.list.return_value.execute.side_effect = [
        {"files": [{"id": "folder_123", "name": "Finance"}]},  # Folder list
        {"files": [{"id": "file_456", "name": "StockMan_Portfolio.xlsx"}]},  # File list
    ]
    mock_files.update.return_value.execute.return_value = {
        "id": "file_456",
        "name": "StockMan_Portfolio.xlsx",
        "modifiedTime": "2026-10-23T10:00:00Z",
    }

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp_path = tmp.name
    tmp.close()

    try:
        res = drive_sync.upload_portfolio_to_drive(mock_service, tmp_path, folder_name="Finance")
        assert res["id"] == "file_456"
        assert res["status"] == "updated"
        assert mock_files.update.called
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def test_download_latest_pick_of_the_day_raw_markdown_native_media():
    """Verify download_latest_pick_of_the_day downloads raw text/markdown natively with alt=media."""
    mock_service = MagicMock()
    mock_files = mock_service.files.return_value

    # Folder list returns 'folder_123', files list returns Pick_of_the_day markdown file
    raw_markdown = (
        "# 🏆 Pick of the Day Report\n\n"
        "| Rank | Code | Score | Action |\n"
        "| :--- | :--- | :---- | :----- |\n"
        "| 1 | HDFC | 92 | ADD |\n"
    )

    mock_files.list.return_value.execute.side_effect = [
        {"files": [{"id": "folder_123", "name": "Finance"}]},
        {
            "files": [
                {
                    "id": "file_markdown_789",
                    "name": "Pick_of_the_day_20261023.md",
                    "mimeType": "text/markdown",
                    "createdTime": "2026-10-23T08:00:00Z",
                    "modifiedTime": "2026-10-23T08:30:00Z",
                }
            ]
        },
    ]

    with patch("StockMan.drive_sync.MediaIoBaseDownload") as mock_downloader_cls:
        mock_instance = MagicMock()
        # Simulate downloader writing content to stream
        def fake_next_chunk():
            return None, True

        mock_instance.next_chunk.side_effect = fake_next_chunk
        mock_downloader_cls.return_value = mock_instance

        with tempfile.TemporaryDirectory() as temp_dir:
            # Inject fake content into the buffer stream
            def mock_init(fh, req):
                fh.write(raw_markdown.encode("utf-8"))
                return mock_instance

            mock_downloader_cls.side_effect = mock_init

            res = drive_sync.download_latest_pick_of_the_day(
                mock_service, folder_name="Finance", dest_dir=temp_dir
            )
            assert res is not None
            saved_path, md_filename = res
            assert md_filename == "Pick_of_the_day_20261023.md"
            assert os.path.exists(saved_path)

            with open(saved_path, "r", encoding="utf-8") as f:
                content = f.read()

            # Verify markdown table pipes are intact
            assert "| Rank | Code | Score | Action |" in content
            assert "| 1 | HDFC | 92 | ADD |" in content

            # Check standard alias file was also saved as .md
            alias_path = os.path.join(temp_dir, "Pick_of_the_day_latest.md")
            assert os.path.exists(alias_path)
            with open(alias_path, "r", encoding="utf-8") as f_alias:
                assert "| 1 | HDFC | 92 | ADD |" in f_alias.read()


def test_download_google_doc_exports_as_markdown():
    """Verify that Google Docs (application/vnd.google-apps.document) are exported with mimeType=text/markdown."""
    mock_service = MagicMock()
    mock_files = mock_service.files.return_value

    mock_doc_content = (
        "# 🏆 Google Doc Pick of the Day\n\n"
        "| Sector | Weight % |\n"
        "| :----- | :------- |\n"
        "| Defense | 39.50% |\n"
    ).encode("utf-8")

    mock_files.list.return_value.execute.side_effect = [
        {"files": [{"id": "folder_123", "name": "Finance"}]},
        {
            "files": [
                {
                    "id": "doc_pick_999",
                    "name": "Pick_of_the_day_Doc.md",
                    "mimeType": "application/vnd.google-apps.document",
                }
            ]
        },
    ]

    mock_export = mock_files.export_media.return_value
    mock_export.execute.return_value = mock_doc_content

    with tempfile.TemporaryDirectory() as temp_dir:
        res = drive_sync.download_latest_pick_of_the_day(
            mock_service, folder_name="Finance", dest_dir=temp_dir
        )
        assert res is not None
        saved_path, md_name = res
        assert md_name == "Pick_of_the_day_Doc.md"
        assert os.path.exists(saved_path)

        with open(saved_path, "r", encoding="utf-8") as f:
            content = f.read()

        assert "| Sector | Weight % |" in content
        assert "| Defense | 39.50% |" in content
        mock_files.export_media.assert_called_with(
            fileId="doc_pick_999", mimeType="text/markdown"
        )


def test_get_current_day_cycle_start():
    """Verify 8:00 AM cycle boundary calculations."""
    from datetime import datetime

    # Scenario A: Exactly at 8:00 AM
    dt_8am = datetime(2026, 10, 24, 8, 0, 0)
    cycle = drive_sync.get_current_day_cycle_start(dt_8am)
    assert cycle == datetime(2026, 10, 24, 8, 0, 0)

    # Scenario B: Afternoon (2:30 PM)
    dt_afternoon = datetime(2026, 10, 24, 14, 30, 0)
    cycle = drive_sync.get_current_day_cycle_start(dt_afternoon)
    assert cycle == datetime(2026, 10, 24, 8, 0, 0)

    # Scenario C: Early Morning (7:59 AM) -> belongs to yesterday's 8:00 AM cycle
    dt_early = datetime(2026, 10, 24, 7, 59, 59)
    cycle = drive_sync.get_current_day_cycle_start(dt_early)
    assert cycle == datetime(2026, 10, 23, 8, 0, 0)


def test_should_prompt_user_today_logic():
    """Verify once-a-day (8:00 AM cycle) prompt check."""
    import json
    from datetime import datetime

    with tempfile.TemporaryDirectory() as temp_dir:
        state_file = os.path.join(temp_dir, "test_state.json")
        fake_excel = os.path.join(temp_dir, "fake.xlsx")

        # 1. No state file, no excel -> should prompt
        now_morning = datetime(2026, 10, 24, 9, 0, 0)
        assert drive_sync.should_prompt_user_today(state_file=state_file, excel_path=fake_excel, now=now_morning) is True

        # 2. Record completion at 9:05 AM on 2026-10-24
        completed_at = datetime(2026, 10, 24, 9, 5, 0)
        drive_sync.record_startup_prompt_completed(
            1800000.0, 150000.0, state_file=state_file, timestamp=completed_at
        )

        # 3. Same day at 2:00 PM -> should NOT prompt
        now_afternoon = datetime(2026, 10, 24, 14, 0, 0)
        assert drive_sync.should_prompt_user_today(state_file=state_file, excel_path=fake_excel, now=now_afternoon) is False

        # 4. Next day before 8:00 AM (e.g. 7:30 AM) -> same cycle -> should NOT prompt
        next_early = datetime(2026, 10, 25, 7, 30, 0)
        assert drive_sync.should_prompt_user_today(state_file=state_file, excel_path=fake_excel, now=next_early) is False

        # 5. Next day after 8:00 AM (e.g. 8:01 AM) -> new cycle -> SHOULD prompt
        next_cycle = datetime(2026, 10, 25, 8, 1, 0)
        assert drive_sync.should_prompt_user_today(state_file=state_file, excel_path=fake_excel, now=next_cycle) is True
